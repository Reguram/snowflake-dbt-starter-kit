#!/usr/bin/env python3
"""
Script to update Snowflake_dbt_Starter_Kit_Overview.pptx with current project context.
Run from the project root with the virtual environment activated.
"""

import copy
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor


def get_shape_text(shape):
    """Return full text of a shape."""
    if shape.has_text_frame:
        return "\n".join(p.text for p in shape.text_frame.paragraphs)
    return ""


def update_text_in_shape(shape, old_text, new_text):
    """Replace old_text with new_text in a shape, preserving run formatting."""
    if not shape.has_text_frame:
        return False

    tf = shape.text_frame
    full_text = tf.text

    if old_text not in full_text:
        return False

    # Find the paragraph and run containing the text
    for para in tf.paragraphs:
        para_text = para.text
        if old_text in para_text:
            # Try to update within individual runs first
            for run in para.runs:
                if old_text in run.text:
                    run.text = run.text.replace(old_text, new_text)
                    return True

            # If it spans runs, rebuild the paragraph text
            # Combine all runs into one with the first run's formatting
            if para.runs:
                first_run = para.runs[0]
                combined = para_text.replace(old_text, new_text)
                first_run.text = combined
                # Clear remaining runs
                for run in para.runs[1:]:
                    run.text = ""
                return True

    return False


def replace_shape_full_text(shape, paragraphs_list):
    """
    Replace entire text frame content with a list of (text, bold, font_size) tuples.
    Each tuple: (text_string, is_bold, font_size_pt or None to keep default)
    """
    if not shape.has_text_frame:
        return

    tf = shape.text_frame
    # Clear existing paragraphs beyond the first
    while len(tf.paragraphs) > 1:
        p = tf.paragraphs[-1]._p
        p.getparent().remove(p)

    # Use existing paragraphs or add new ones
    for i, (text, bold, font_size) in enumerate(paragraphs_list):
        if i == 0:
            para = tf.paragraphs[0]
        else:
            para = tf.add_paragraph()

        # Clear existing runs
        for run in para.runs:
            run.text = ""

        if para.runs:
            run = para.runs[0]
        else:
            run = para.add_run()

        run.text = text
        if bold is not None:
            run.font.bold = bold
        if font_size is not None:
            run.font.size = Pt(font_size)


def main():
    pptx_path = "Snowflake_dbt_Starter_Kit_Overview.pptx"
    prs = Presentation(pptx_path)

    slides = prs.slides
    changes_made = []

    # =========================================================================
    # SLIDE 3: Key Features
    # =========================================================================
    slide3 = slides[2]
    for shape in slide3.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text

        # Update MCP Server tool count
        if "11+ tools for Claude/Copilot: model generation, code review, data quality" in text:
            updated = update_text_in_shape(
                shape,
                "11+ tools for Claude/Copilot: model generation, code review, data quality",
                "6 local / 8 managed tools: model gen, code review, data quality, Cortex Analyst"
            )
            if updated:
                changes_made.append("Slide 3: Updated MCP Server tool count")

        # Update AI Agent Skills count
        if "11 dbt-labs skills + custom skills for VS Code Copilot auto-activation" in text:
            updated = update_text_in_shape(
                shape,
                "11 dbt-labs skills + custom skills for VS Code Copilot auto-activation",
                "50 total skills (17 dbt/custom + 33 bundled Snowflake) for Claude Code, Copilot & Cortex"
            )
            if updated:
                changes_made.append("Slide 3: Updated AI Agent Skills count")

        # Update Evaluation Framework description
        if "8 tasks to benchmark Cortex Code vs Claude Code vs Copilot CLI" in text:
            updated = update_text_in_shape(
                shape,
                "8 tasks to benchmark Cortex Code vs Claude Code vs Copilot CLI",
                "8 tasks benchmarking Cortex Code vs Claude Code vs Copilot CLI across 8 weighted dimensions"
            )
            if updated:
                changes_made.append("Slide 3: Updated Evaluation Framework description")

    # =========================================================================
    # SLIDE 4: Architecture Overview
    # =========================================================================
    slide4 = slides[3]
    for shape in slide4.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text

        if "AI-powered tools (11+) running as Snowflake UDFs" in text:
            updated = update_text_in_shape(
                shape,
                "AI-powered tools (11+) running as Snowflake UDFs",
                "AI-powered tools running as Snowflake UDFs (8 managed)"
            )
            if updated:
                changes_made.append("Slide 4: Updated MCP Server tool count in architecture")

        if "or local Python server for model gen, review, quality" in text:
            updated = update_text_in_shape(
                shape,
                "or local Python server for model gen, review, quality",
                "or local Python server (6 tools) for model gen, review, quality"
            )
            if updated:
                changes_made.append("Slide 4: Updated local MCP server tool count")

        if "11 dbt-labs skills + custom Snowflake skills" in text:
            updated = update_text_in_shape(
                shape,
                "11 dbt-labs skills + custom Snowflake skills",
                "50 total skills (17 dbt/custom + 33 bundled Snowflake)"
            )
            if updated:
                changes_made.append("Slide 4: Updated AI Agent Skills count in architecture")

        if "Auto-activate in VS Code on file open" in text:
            updated = update_text_in_shape(
                shape,
                "Auto-activate in VS Code on file open",
                "Active in Claude Code, GitHub Copilot & Cortex Code"
            )
            if updated:
                changes_made.append("Slide 4: Updated skills activation description")

    # =========================================================================
    # SLIDE 6: Technology Stack
    # =========================================================================
    slide6 = slides[5]
    for shape in slide6.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text

        # Update dbt Agent Skills description
        if "11 upstream skills from dbt-labs" in text:
            updated = update_text_in_shape(
                shape,
                "11 upstream skills from dbt-labs",
                "50 total skills: 17 dbt/custom + 33 bundled Snowflake skills"
            )
            if updated:
                changes_made.append("Slide 6: Updated dbt Agent Skills count in tech stack")

        # Update GitHub Copilot label to reflect all AI tools
        if shape.text_frame.text.strip() == "GitHub Copilot":
            updated = update_text_in_shape(
                shape,
                "GitHub Copilot",
                "Claude Code / Copilot"
            )
            if updated:
                changes_made.append("Slide 6: Updated GitHub Copilot label to Claude Code / Copilot")

        if "AI pair programming with custom skills & agents" in text:
            updated = update_text_in_shape(
                shape,
                "AI pair programming with custom skills & agents",
                "AI pair programming with 50 skills & custom agents"
            )
            if updated:
                changes_made.append("Slide 6: Updated Copilot description")

    # =========================================================================
    # SLIDE 7: Folder / Module Structure
    # =========================================================================
    slide7 = slides[6]
    for shape in slide7.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text

        # Update MCP server tool count
        if "server.py  —  Local fallback MCP server (11+ tools)" in text:
            updated = update_text_in_shape(
                shape,
                "server.py  —  Local fallback MCP server (11+ tools)",
                "server.py  —  Local fallback MCP server (6 tools)"
            )
            if updated:
                changes_made.append("Slide 7: Updated local MCP server tool count")

        # Update MCP setup scripts description
        if "setup/  —  Snowflake Managed MCP setup SQL (5 scripts)" in text:
            updated = update_text_in_shape(
                shape,
                "setup/  —  Snowflake Managed MCP setup SQL (5 scripts)",
                "setup/  —  Snowflake Managed MCP setup SQL (5 scripts, 8 tools)"
            )
            if updated:
                changes_made.append("Slide 7: Added managed MCP tool count")

        # Update evaluation REPORT.md description
        if "REPORT.md  —  Cortex vs Claude vs Copilot results" in text:
            updated = update_text_in_shape(
                shape,
                "REPORT.md  —  Cortex vs Claude vs Copilot results",
                "REPORT.md  —  Cortex vs Claude vs Copilot results template (8 dimensions)"
            )
            if updated:
                changes_made.append("Slide 7: Updated evaluation REPORT.md description")

    # =========================================================================
    # SLIDE 8: Key Integrations
    # =========================================================================
    slide8 = slides[7]
    for shape in slide8.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text

        # Update GitHub Copilot / VS Code section
        if "11 auto-activating skills, 4 reusable prompts, and a custom @dbt-semantic-advisor agent" in text:
            updated = update_text_in_shape(
                shape,
                "11 auto-activating skills, 4 reusable prompts, and a custom @dbt-semantic-advisor agent. Skills trigger contextually based on file type.",
                "50 skills (17 dbt/custom + 33 bundled Snowflake), reusable prompts, and a custom @dbt-semantic-advisor agent. Skills activate contextually in Claude Code, Copilot & Cortex."
            )
            if updated:
                changes_made.append("Slide 8: Updated Copilot skills count and tools")
            else:
                # Try partial update
                updated = update_text_in_shape(
                    shape,
                    "11 auto-activating skills, 4 reusable prompts",
                    "50 auto-activating skills (17 dbt/custom + 33 bundled Snowflake)"
                )
                if updated:
                    changes_made.append("Slide 8: Updated skills count (partial)")

        # Update dbt-labs Agent Skills section
        if "11 open-source skills from dbt-labs covering analytics engineering" in text:
            updated = update_text_in_shape(
                shape,
                "11 open-source skills from dbt-labs covering analytics engineering, semantic layer, unit testing, CLI operations, troubleshooting, and cross-platform migration.",
                "50 total skills: 17 covering dbt analytics engineering, semantic layer, unit testing, CLI ops & migration — plus 33 bundled Snowflake skills spanning Cortex AI, dynamic tables, ML, Iceberg, Streamlit, data governance, and more."
            )
            if updated:
                changes_made.append("Slide 8: Updated dbt-labs Agent Skills description")

        # Update dbt-labs Agent Skills label
        if shape.text_frame.text.strip() == "dbt-labs Agent Skills":
            updated = update_text_in_shape(
                shape,
                "dbt-labs Agent Skills",
                "AI Agent Skills (50 Total)"
            )
            if updated:
                changes_made.append("Slide 8: Updated dbt-labs Agent Skills label")

    # Save the updated file
    output_path = "Snowflake_dbt_Starter_Kit_Overview.pptx"
    prs.save(output_path)

    print(f"Saved updated presentation to: {output_path}")
    print(f"\nChanges made ({len(changes_made)}):")
    for change in changes_made:
        print(f"  ✓ {change}")

    if not changes_made:
        print("  No changes were made. Text may have already been updated or not found.")


if __name__ == "__main__":
    main()
